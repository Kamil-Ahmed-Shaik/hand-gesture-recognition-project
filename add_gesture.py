import time
from openpyxl import load_workbook, Workbook
import cv2
import mediapipe as mp
import os

cap = cv2.VideoCapture(0)

mphands = mp.solutions.hands
hands = mphands.Hands(max_num_hands=1)  # Limit to 1 hand for performance optimization
mpdraw = mp.solutions.drawing_utils
ptime = 0
ctime = 0

path = "Data_source.xlsx"

def dist(x, y):
    return ((x[1] - y[1]) ** 2 + (x[2] - y[2]) ** 2) ** 0.5 / 1000

def fun(a):
    ini = a[0]
    dis = a[1:]
    result = []

    for i in dis:
        if i[0] not in [1, 2, 5, 9, 13, 17]:
            dip = dist(i, ini)
            result.append(dip)

    lst = [4, 8, 12, 16, 20]
    
    for i in range(len(lst) - 1):
        tp = []
        for j in range(i + 1, len(lst)):
            tp.append(dist(a[lst[i]], a[lst[j]]))
        result.extend(tp)
    
    return result

atlast = []
new_gest = input("Enter the new gesture: ")
print("Press 't' to capture the gesture\nPress 'q' to discard the details\nPress 'f' to save the gesture details")

# Check if the Excel file exists, create it if it doesn't
if not os.path.exists(path):
    wb = Workbook()
    wb.save(path)  # Just create the file without adding headers

while True:
    success, img = cap.read()
    if not success:
        continue

    ky = cv2.waitKey(1)
    imgRGB = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
    results = hands.process(imgRGB)

    if ky == ord('t') and results.multi_hand_landmarks:
        for handLms in results.multi_hand_landmarks:
            temp = [[id, int(lm.x * img.shape[1]), int(lm.y * img.shape[0])] for id, lm in enumerate(handLms.landmark)]
            print(temp)

            tp = fun(temp)
            tp.append(new_gest)
            atlast.append(tp)

        mpdraw.draw_landmarks(img, handLms, mphands.HAND_CONNECTIONS)

    ctime = time.time()
    if ctime - ptime > 0:  # Prevent division by zero
        fps = 1 / (ctime - ptime)
    else:
        fps = 0
    ptime = ctime

    cv2.putText(img, str(int(fps)), (10, 70), cv2.FONT_HERSHEY_PLAIN, 3, (0, 0, 255), 3)
    cv2.imshow("Image", img)

    if ky == ord('q'):
        break
    elif ky == ord('f'):
        try:
            wb = load_workbook(path)
            ws = wb.active
        except FileNotFoundError:
            print(f"Error: The file '{path}' was not found.")
            break

        for row in atlast:
            ws.append(row)

        wb.save(path)
        break

print(atlast)
cv2.destroyAllWindows()
cap.release()
